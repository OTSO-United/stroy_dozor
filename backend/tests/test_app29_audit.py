"""APP-29: behavior at temporal, lifecycle, evidence and export boundaries."""

import csv
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
import numpy as np
from openpyxl import load_workbook
from sqlalchemy import select

from app import models as m, monitor, schemas, security
from app.jobs import claim
from app.telemetry import series_points
from app.temporal_cv import ActivityAnalyzer, LocalTracker
from test_app import post, work

START = datetime(2026, 9, 28, 6, tzinfo=timezone.utc)


@pytest.mark.parametrize("scenario", ["body", "mechanism", "idle", "insufficient"])
def test_activity_from_synthetic_pixels_and_tracks(scenario):
    frame = np.random.default_rng(29).integers(30, 220, (320, 320, 3), dtype=np.uint8)
    tracker = LocalTracker(max_gap=3)
    analyzer = ActivityAnalyzer(sample_seconds=1, idle_after=10, method="difference")
    for tick in range(25):
        time = tick / 2
        x = 0.1 + tick * 0.012 if scenario == "body" else 0.2
        current = frame.copy()
        if scenario == "mechanism":
            current[100:140, 85:110] = 40 if tick % 2 else 240
        if scenario == "insufficient":
            current[:] = 0
        if tick % 2 == 0:
            detections = tracker.update(
                [{"class_id": "1", "confidence": 0.9, "bbox": [x, 0.2, x + 0.2, 0.7]}],
                time,
                current,
            )
        analyzer.observe(current, time, tracker.tracks)
        final = analyzer.annotate(detections, time, tracker.tracks)[0]
    expected = {
        "body": ("working", "vehicle_motion"),
        "mechanism": ("working", "visible_mechanism_motion"),
        "idle": ("idle", "visible_inactivity"),
        "insufficient": ("unknown", "camera_registration_uncertain"),
    }
    assert (final["activity"], final["activity_basis"]) == expected[scenario]


def samples(levels, *, step=1):
    return [
        SimpleNamespace(
            id=str(i),
            run_id="run",
            source_id="camera",
            captured_at=START + timedelta(seconds=i * step),
            offset_seconds=i * step,
            quality={"usable": value is not None},
            model={"supported_classes": ["0"]},
            detections=[{"class_id": "0", "activity": "unknown"}] * (value or 0),
        )
        for i, value in enumerate(levels)
    ]


@pytest.mark.parametrize("stable,temporary", [(3, 4), (3, 2), (3, 0), (0, 3)])
def test_brief_changes_do_not_change_stable_line(stable, temporary):
    rows = samples([stable] * 35 + [temporary] * 9 + [stable] * 16)
    points = series_points(rows, [0], 0, 3)
    assert all(point["counts"]["0"] == stable for point in points[34:])
    [minute] = series_points(rows, [0], 60, 3)
    assert minute["counts"]["0"] == stable
    assert rows[35].detections == samples([temporary])[0].detections


@pytest.mark.parametrize(
    "before,after,threshold", [(3, 4, 10), (3, 2, 10), (3, 0, 30), (0, 2, 10)]
)
def test_change_is_confirmed_at_exact_event_time_boundary(before, after, threshold):
    points = series_points(
        samples([before] * 40 + [after] * (threshold + 2)), [0], 0, 3
    )
    assert points[40 + threshold - 1]["counts"]["0"] == before
    assert points[40 + threshold]["counts"]["0"] == after


def test_minute_does_not_restore_repeated_short_raw_spikes():
    levels = [3] * 12 + ([4] * 4 + [3] * 4) * 6 + [2] * 60
    points = series_points(samples(levels), [0], 60, 3)
    assert [point["counts"]["0"] for point in points] == [3, 2]


def test_gap_unknown_source_and_run_each_reset_confirmation():
    for field, value in [("run_id", "next"), ("source_id", "other")]:
        rows = samples([3, 4, 4])
        setattr(rows[2], field, value)
        points = series_points(rows, [0], 0, 3)
        assert points[2]["segment"] != points[1]["segment"]
    rows = samples([3, 4, None, 4])
    points = series_points(rows, [0], 0, 3)
    assert points[2]["counts"]["0"] is None
    assert points[3]["segment"] != points[1]["segment"]
    rows = samples([3, 4, 4])
    rows[2].captured_at += timedelta(seconds=10)
    assert series_points(rows, [0], 0, 3)[2]["segment"] == 1


def test_activity_unknown_and_measured_zero_remain_distinct():
    points = series_points(samples([2] * 60), [0], 60, 3)
    assert points[0]["display_activity"] is None
    assert points[0]["display_activity_by_class"]["0"] == {
        "working": None,
        "idle": None,
    }
    rows = samples([0] * 60)
    assert series_points(rows, [0], 60, 3)[0]["display_activity"] == {
        "working": 0,
        "idle": 0,
    }
    rows = samples([1, 1])
    rows[0].detections = [
        {"class_id": "0", "activity": "working", "activity_basis": "vehicle_motion"}
    ]
    rows[1].detections = [
        {"class_id": "0", "activity": "idle", "activity_basis": "visible_inactivity"}
    ]
    assert series_points(rows, [0], 60, 3)[0]["display_activity"] == {
        "working": 0,
        "idle": 1,
    }


def test_presence_confirmation_does_not_hide_classified_activity():
    rows = samples([1, 2, 2])
    for row in rows:
        row.detections = [
            {"class_id": "0", "activity": "working"} for _ in row.detections
        ]
    points = series_points(rows, [0], 0, 3)
    assert points[-1]["counts"]["0"] == 1
    assert points[-1]["display_activity_by_class"]["0"]["working"] == 2


def test_sparse_minutes_use_historical_sampling_and_preserve_real_gaps():
    rows = samples([2, 2, 2, 2], step=60)
    rows[-1].captured_at += timedelta(minutes=2)
    points = series_points(rows, [0], 60, 3, expected_step=1, steps_by_run={"run": 60})
    assert [point["segment"] for point in points] == [0, 0, 0, 1]
    assert all(point["counts"]["0"] == 2 for point in points)
    assert all(point["display_basis"]["0"] == "sampled" for point in points)
    assert all(
        point["activity_seconds_by_class"]["0"]["working"] is None for point in points
    )


def test_machine_time_uses_raw_classified_intervals_before_aggregation():
    rows = samples([1] * 120)
    for index, row in enumerate(rows):
        row.detections = [
            {
                "class_id": "0",
                "activity": "working"
                if index < 30
                else "idle"
                if index < 90
                else "unknown",
            }
        ]
    for step in (0, 60, 120):
        points = series_points(rows, [0], step, 3, expected_step=1)
        totals = {
            state: sum(
                point["activity_seconds_by_class"]["0"][state] or 0 for point in points
            )
            for state in ("working", "idle")
        }
        assert totals == {"working": 30, "idle": 60}
        assert (
            sum(point["presence_seconds_by_class"]["0"]["present"] for point in points)
            == 119
        )
    window = series_points(
        rows, [0], 60, 3, expected_step=1, display_from=START + timedelta(seconds=40)
    )
    assert (
        sum(point["activity_seconds_by_class"]["0"]["working"] or 0 for point in window)
        == 0
    )
    assert (
        sum(point["activity_seconds_by_class"]["0"]["idle"] or 0 for point in window)
        == 50
    )


@pytest.fixture
def timeline(runtime, project):
    factory, folder = runtime
    stage = work(
        start=START.isoformat(),
        end=(START + timedelta(hours=4)).isoformat(),
        resources={"0": 2, "1": 1, "16": 0},
    )
    with factory() as db:
        plan = m.Plan(project_id=project["id"], version=1, works=[stage])
        source = m.Source(
            project_id=project["id"],
            name="Камера аудита",
            kind="video",
            sample_seconds=1,
            metadata_json={"width": 320, "height": 200},
        )
        db.add_all([plan, source])
        db.flush()
        db.get(m.Project, project["id"]).current_plan_id = plan.id
        binding = m.Binding(
            project_id=project["id"],
            source_id=source.id,
            revision=1,
            regions=[
                {
                    "id": "whole",
                    "name": "Весь кадр",
                    "work_ids": [stage["id"]],
                    "polygon": [[0, 0], [1, 0], [1, 1], [0, 1]],
                    "primary": True,
                    "visibility_confirmed": True,
                }
            ],
        )
        db.add(binding)
        db.flush()
        settings = schemas.Settings().model_dump()
        job = m.Job(
            project_id=project["id"],
            source_id=source.id,
            kind="analyze",
            payload={
                "plan_id": plan.id,
                "binding_id": binding.id,
                "settings": settings,
                "sample_seconds": 1,
                "works": [{k: stage[k] for k in ("id", "code", "title")}],
                "work_states": {},
            },
        )
        db.add(job)
        db.commit()
        ids = {
            "project": project["id"],
            "source": source.id,
            "job": job.id,
            "plan": plan.id,
            "binding": binding.id,
            "stage": stage,
        }

    def add(levels, *, start=0, gap=1, activity="unknown"):
        with factory() as db:
            for i, count in enumerate(levels, start):
                detections = [
                    {
                        "class_id": "0",
                        "confidence": 0.9,
                        "bbox": [0.2, 0.2, 0.4, 0.8],
                        "activity": activity,
                    }
                ] * count
                detections += [
                    {
                        "class_id": "1",
                        "confidence": 0.9,
                        "bbox": [0.5, 0.2, 0.7, 0.8],
                        "activity": "unknown",
                    }
                ]
                db.add(
                    m.Observation(
                        project_id=ids["project"],
                        source_id=ids["source"],
                        run_id=ids["job"],
                        sample_index=i,
                        captured_at=START + timedelta(seconds=i * gap),
                        offset_seconds=i * gap,
                        created_at=START + timedelta(seconds=i),
                        detections=detections,
                        quality={"usable": True},
                        model={"supported_classes": ["0", "1"]},
                        evidence_key=f"evidence/audit-{i}.jpg",
                    )
                )
            db.commit()

    return ids, add, factory, folder


@pytest.mark.parametrize("step", [0, 60])
def test_api_pages_and_time_window_do_not_reset_count_state(client, timeline, step):
    ids, add, _, _ = timeline
    add([3] * 55 + [4] * 9 + [3] * 9 + [2] * 57)
    url = f"/api/v1/projects/{ids['project']}/timeseries"
    base = {
        "source_id": ids["source"],
        "work_id": ids["stage"]["id"],
        "display_seconds": step,
    }
    full = client.get(url, params={**base, "limit": 5000}).json()["points"]
    cursor, pages = {}, []
    for _ in range(150):
        page = client.get(url, params={**base, "limit": 7, **cursor}).json()
        pages = page["points"] + pages
        if not page["truncated"]:
            break
        cursor = {
            "before": page["next_before"],
            "before_id": page["next_before_id"],
            "before_sample_index": page["next_before_sample_index"],
        }
    assert pages == full
    assert len({point["id"] for point in pages}) == len(pages)
    if step == 0:
        window = client.get(
            url, params={**base, "from": (START + timedelta(seconds=58)).isoformat()}
        ).json()["points"]
        assert window == full[58:]
        assert full[60]["counts"]["0"] == 3
        assert full[83]["counts"]["0"] == 2


def test_stage_timeseries_includes_detected_class_absent_from_plan(client, timeline):
    ids, add, factory, _ = timeline
    add([2] * 20)
    with factory() as db:
        observation = db.scalar(
            select(m.Observation)
            .where(m.Observation.run_id == ids["job"])
            .order_by(m.Observation.sample_index)
        )
        observation.model = {"supported_classes": ["0", "1", "2"]}
        observation.detections = [
            *observation.detections,
            {"class_id": "2", "confidence": 0.9, "bbox": [0.2, 0.2, 0.4, 0.8]},
        ]
        db.commit()
    data = client.get(
        f"/api/v1/projects/{ids['project']}/timeseries",
        params={"source_id": ids["source"], "work_id": ids["stage"]["id"]},
    ).json()
    assert any(point["counts"].get("2") == 1 for point in data["points"])


def test_job_readiness_waits_for_every_assessment_and_preserves_snapshot(
    client, timeline
):
    ids, add, factory, _ = timeline
    add([1, 1])
    url = f"/api/v1/projects/{ids['project']}/jobs"
    assert client.get(url).json()[0]["analysis_state"] == "queued"
    with factory() as db:
        started = claim(db)
        assert started.payload["started_at"]
    assert client.get(url).json()[0]["analysis_state"] == "running"
    with factory() as db:
        db.get(m.Job, ids["job"]).status = "succeeded"
        db.get(m.Binding, ids["binding"]).regions = []
        db.commit()
    assert client.get(url).json()[0]["analysis_state"] == "assessing"
    monitor.tick(limit=1, project_id=ids["project"])
    assert not client.get(url).json()[0]["results_ready"]
    monitor.tick(project_id=ids["project"])
    result = client.get(url).json()[0]
    assert result["analysis_state"] == "ready"
    assert result["works"][0]["code"] == ids["stage"]["code"]


@pytest.mark.parametrize("kind", ["video", "image", "rtsp"])
def test_ready_bound_source_queues_once_with_snapshot(
    client, timeline, monkeypatch, kind
):
    from app import api

    ids, _, factory, _ = timeline
    monkeypatch.setattr(api, "model_status", lambda: {"ready": True, "sha256": "audit"})
    with factory() as db:
        source = db.get(m.Source, ids["source"])
        source.kind = kind
        source.metadata_json = {}
        db.get(m.Job, ids["job"]).status = "succeeded"
        db.commit()
    url = f"/sources/{ids['source']}/bindings"
    current = client.get("/api/v1" + url).json()
    first = post(
        client,
        url,
        {"expected_revision": current["revision"], "regions": current["regions"]},
    )
    assert first.status_code == 200
    jobs_url = f"/api/v1/projects/{ids['project']}/jobs"
    assert len(client.get(jobs_url).json()) == 1
    with factory() as db:
        db.get(m.Source, ids["source"]).metadata_json = {"width": 320, "height": 200}
        db.commit()
    for _ in range(2):
        current = client.get("/api/v1" + url).json()
        assert (
            post(
                client,
                url,
                {
                    "expected_revision": current["revision"],
                    "regions": current["regions"],
                },
            ).status_code
            == 200
        )
    queued = [job for job in client.get(jobs_url).json() if job["status"] == "queued"]
    assert len(queued) == 1
    assert queued[0]["works"][0]["id"] == ids["stage"]["id"]
    assert queued[0]["started_at"] is None
    assert queued[0]["continuous"] == (kind == "rtsp")


@pytest.mark.parametrize("decision", ["confirmed", "dismissed"])
def test_review_and_new_episode_are_independent_and_audited(client, timeline, decision):
    ids, add, factory, _ = timeline
    add([0] * 32)
    monitor.tick(project_id=ids["project"])
    url = f"/api/v1/projects/{ids['project']}/alerts"
    first = client.get(url).json()[0]
    assert first["status"] == "open"
    saved = post(
        client,
        f"/alerts/{first['id']}/reviews",
        {"decision": decision, "note": "Проверка"},
    )
    assert saved.status_code == 200
    assert saved.json()["data"]["decision"] == decision
    assert (
        client.get(f"/api/v1/projects/{ids['project']}").json()["summary"][
            "open_alerts"
        ]
        == 0
    )
    add([2] + [0] * 32, start=32)
    monitor.tick(project_id=ids["project"])
    alerts = client.get(url).json()
    assert len(alerts) == 2
    old = next(alert for alert in alerts if alert["id"] == first["id"])
    new = next(alert for alert in alerts if alert["id"] != first["id"])
    assert old["reviewed"] and old["details"]["review"]["decision"] == decision
    assert new["status"] == "open" and not new["reviewed"]
    assert new["details"]["episode"] == 2
    assert (
        client.get(f"/api/v1/projects/{ids['project']}").json()["summary"][
            "open_alerts"
        ]
        == 1
    )
    with factory() as db:
        assert not list(db.scalars(select(m.WorkState)))


def test_evidence_only_exposes_confirmed_missing_classes(client, timeline):
    ids, add, _, _ = timeline
    add([0] * 35)
    monitor.tick(project_id=ids["project"])
    alert = client.get(f"/api/v1/projects/{ids['project']}/alerts").json()[0]
    evidence = client.get(
        f"/api/v1/alerts/{alert['id']}/evidence", params={"class_id": "0"}
    ).json()
    assert evidence["missing_class_ids"] == ["0"]
    assert evidence["keyframes"] and evidence["frames"]
    assert all(
        frame["offset_seconds"] >= 30 and frame["run_id"] == ids["job"]
        for frame in evidence["frames"]
    )
    assert [stat["class_id"] for stat in evidence["absence_stats"]] == ["0"]



def test_single_class_alert_keeps_keyframe_when_episode_history_is_short(client, timeline):
    ids, add, factory, _ = timeline
    add([0] * 35)
    monitor.tick(project_id=ids["project"])
    alert = client.get(f"/api/v1/projects/{ids['project']}/alerts").json()[0]
    with factory() as db:
        saved = db.get(m.Alert, alert["id"])
        saved.first_seen = saved.last_seen
        db.commit()
    evidence = client.get(f"/api/v1/alerts/{alert['id']}/evidence").json()
    assert evidence["deviation_class_ids"] == ["0"]
    assert evidence["keyframes"]
    assert evidence["basis"]["id"] in {frame["id"] for frame in evidence["keyframes"]}


@pytest.mark.parametrize("kind", ["missing", "excess", "idle"])
def test_evidence_classes_follow_current_alert_not_other_episode_frames(
    client, timeline, kind
):
    ids, add, factory, _ = timeline
    add([2] * 70)
    with factory() as db:
        rows = list(
            db.scalars(
                select(m.Observation)
                .where(m.Observation.run_id == ids["job"])
                .order_by(m.Observation.sample_index)
            )
        )
        for index, row in enumerate(rows):
            detections = [dict(item) for item in row.detections]
            if kind == "missing":
                detections = [
                    item
                    for item in detections
                    if item["class_id"] != ("1" if index < 35 else "0")
                ]
            elif kind == "excess":
                extra = dict(
                    next(
                        item
                        for item in detections
                        if item["class_id"] == ("1" if index < 35 else "0")
                    )
                )
                detections.append(extra)
            else:
                for item in detections:
                    item["activity"] = (
                        "idle"
                        if item["class_id"] == ("1" if index < 35 else "0")
                        else "working"
                    )
            row.detections = detections
        db.commit()
    monitor.tick(project_id=ids["project"])
    alerts = client.get(f"/api/v1/projects/{ids['project']}/alerts").json()
    alert = next(item for item in alerts if item["kind"] == kind)
    evidence = client.get(f"/api/v1/alerts/{alert['id']}/evidence")
    assert evidence.status_code == 200
    data = evidence.json()
    assert data["deviation_class_ids"] == ["0"]
    assert all(
        set(frame["deviation_class_ids"]) <= {"0"}
        for frame in data["keyframes"]
    )
    assert (
        client.get(
            f"/api/v1/alerts/{alert['id']}/evidence",
            params={"class_id": "1"},
        ).status_code
        == 422
    )


def test_old_pending_alerts_are_not_hidden_by_the_recent_history_limit(
    client, timeline
):
    ids, add, factory, _ = timeline
    add([0] * 32)
    monitor.tick(project_id=ids["project"])
    with factory() as db:
        first = db.scalar(select(m.Alert))
        for index in range(1, 211):
            instant = START + timedelta(minutes=index)
            db.add(
                m.Alert(
                    project_id=ids["project"],
                    source_id=ids["source"],
                    run_id=ids["job"],
                    work_id=ids["stage"]["id"],
                    fingerprint=first.fingerprint,
                    kind="missing",
                    severity="warning",
                    status="open",
                    first_seen=instant,
                    last_seen=instant,
                    assessment_id=first.assessment_id,
                    details={**first.details, "episode": index + 1},
                )
            )
        db.commit()
    root = f"/api/v1/projects/{ids['project']}"
    assert len(client.get(root + "/alerts").json()) == 211
    assert len(client.get(root + "/monitoring").json()["alerts"]) == 211
    assert client.get(root).json()["summary"]["open_alerts"] == 211
    listed = next(
        project
        for project in client.get("/api/v1/projects").json()
        if project["id"] == ids["project"]
    )
    assert listed["summary"]["open_alerts"] == 211


def test_xlsx_and_csv_values_historical_plan_period_unknown_and_empty(client, timeline):
    ids, add, factory, _ = timeline
    add([0] * 35)
    monitor.tick(project_id=ids["project"])
    with factory() as db:
        alert = db.scalar(select(m.Alert))
        alert.details = {
            **alert.details,
            "notified_at": None,
            "counts": {"0": 0, "1": None},
        }
        newer = m.Plan(
            project_id=ids["project"],
            version=2,
            works=[
                {**ids["stage"], "starts_at": (START + timedelta(days=1)).isoformat()}
            ],
        )
        db.add(newer)
        db.flush()
        db.get(m.Project, ids["project"]).current_plan_id = newer.id
        db.commit()
    period = {
        "from": (START + timedelta(seconds=5)).isoformat(),
        "to": (START + timedelta(seconds=34)).isoformat(),
    }
    url = f"/api/v1/projects/{ids['project']}/reports/"
    response = client.get(url + "xlsx", params=period)
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    sheet = book["События"]
    headers, values = list(sheet.values)
    row = dict(zip(headers, values))
    assert row["Объект"] == "Тестовая стройка"
    assert row["Индекс этапа"] == ids["stage"]["code"]
    assert row["Начало этапа, МСК"] == "28.09.2026 09:00:00"
    assert row["Начало периода отчёта, МСК"] == "28.09.2026 09:00:05"
    assert row["Конец периода отчёта, МСК"] == "28.09.2026 09:00:34"
    assert row["Начало события, МСК"] == "28.09.2026 09:00:00"
    assert row["Конец события, МСК"] is None
    assert row["Уведомление, МСК"] is None
    assert row["Самосвал: план"] == 2
    assert row["Самосвал: на кадре"] == 0
    assert row["Экскаватор: на кадре"] is None
    csv_rows = list(
        csv.reader(
            StringIO(
                client.get(url + "csv", params=period).content.decode("utf-8-sig")
            ),
            delimiter=";",
        )
    )
    assert csv_rows[0] == list(headers)
    assert csv_rows[1] == ["" if value is None else str(value) for value in values]
    empty = load_workbook(
        BytesIO(
            client.get(
                url + "xlsx", params={"from": (START + timedelta(days=2)).isoformat()}
            ).content
        )
    )
    assert empty["События"].max_row == 1
    assert empty["Период"]["B2"].value == "30.09.2026 09:00:00"
    assert (
        client.get(
            url + "xlsx", params={"from": period["to"], "to": period["from"]}
        ).status_code
        == 422
    )


def test_short_unconfirmed_episode_is_absent_from_report(client, timeline):
    ids, add, _, _ = timeline
    add([0] * 5 + [2])
    monitor.tick(project_id=ids["project"])
    response = client.get(f"/api/v1/projects/{ids['project']}/reports/xlsx")
    assert load_workbook(BytesIO(response.content))["События"].max_row == 1


@pytest.mark.parametrize(
    "value,status", [(-1, 422), (0, 200), (10, 200), (3600, 200), (3601, 422)]
)
def test_count_setting_api_validation_and_persistence(client, project, value, status):
    result = client.put(
        f"/api/v1/projects/{project['id']}/settings",
        json={
            **project["settings"],
            "expected_revision": project["revision"],
            "count_change_confirm_seconds": value,
        },
        headers={"Idempotency-Key": str(uuid4())},
    )
    assert result.status_code == status
    if status == 200:
        settings = client.get(f"/api/v1/projects/{project['id']}").json()["settings"]
        assert settings["count_change_confirm_seconds"] == value
        assert settings["absence_confirm_seconds"] == 30


@pytest.mark.parametrize("scheme", ["rtsp", "rtsps", "https"])
def test_network_source_api_keeps_encrypted_uri_and_probe(
    client, project, runtime, scheme
):
    uri = f"{scheme}://user:password@camera.example/live.m3u8"
    response = post(
        client, f"/projects/{project['id']}/sources", {"name": "Поток", "uri": uri}
    )
    assert response.status_code == 200, response.text
    assert "password" not in response.text
    factory, _ = runtime
    with factory() as db:
        source = db.scalar(select(m.Source))
        assert source.kind == "rtsp" and source.uri_encrypted != uri
        assert security.decode_uri(source.uri_encrypted) == uri
        assert db.scalar(select(m.Job)).kind == "probe"


def test_timeseries_covers_two_hour_video_and_opens_exact_observation(client, timeline):
    ids, add, _, _ = timeline
    add([1] * 13, gap=600)
    points = client.get(
        f"/api/v1/projects/{ids['project']}/timeseries",
        params={
            "source_id": ids["source"],
            "work_id": ids["stage"]["id"],
            "display_seconds": 300,
            "limit": 5000,
        },
    ).json()["points"]
    assert len(points) == 13
    assert points[-1]["offset_seconds"] - points[0]["offset_seconds"] == 7200
    observation = client.get(f"/api/v1/observations/{points[6]['id']}")
    assert observation.status_code == 200
    assert observation.json()["id"] == points[6]["id"]
    assert len(observation.json()["detections"]) == 2
