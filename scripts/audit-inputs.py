"""Read-only audit of organizer inputs. Never extracts archives or writes to sources."""
import argparse
import collections
import hashlib
import json
import re
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from pypdf import PdfReader

def digest(path):
    sha = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            sha.update(block)
    return sha.hexdigest()

def serial(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value

def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", required=True)
    parser.add_argument("--output-root", required=True)
    args = parser.parse_args()
    source = Path(args.source_root).resolve()
    output = Path(args.output_root).resolve()
    data = source / "datasets" / "7.ДГП_датасеты"
    # Sources and reports may share the project root, but never the raw dataset folder.
    dataset_root = (source / "datasets").resolve()
    if output == dataset_root or dataset_root in output.parents:
        raise ValueError("Outputs must not be inside datasets")
    target = output / "docs" / "data"
    scratch = output / "tmp" / "input-audit"
    target.mkdir(parents=True, exist_ok=True)
    scratch.mkdir(parents=True, exist_ok=True)
    report = {"audit_date": "2026-09-15", "source_root": str(source), "files": [], "images": [], "workbooks": [], "archives": []}
    for path in sorted(data.rglob("*")):
        if path.is_file():
            report["files"].append({"path": path.relative_to(source).as_posix(), "bytes": path.stat().st_size})
    images = sorted(data.rglob("*.png"), key=lambda p: int(re.search(r"(\d+)", p.stem).group(1)))
    for path in images:
        with Image.open(path) as img:
            record = {"name": path.name, "width": img.width, "height": img.height, "mode": img.mode,
                      "metadata_keys": sorted(img.info.keys()), "exif_keys": [str(k) for k in img.getexif().keys()],
                      "bytes": path.stat().st_size, "sha256": digest(path)}
            report["images"].append(record)
        with Image.open(path) as verification_image:
            verification_image.verify()
    hashes = collections.defaultdict(list)
    for image in report["images"]:
        hashes[image["sha256"]].append(image["name"])
    report["exact_duplicate_groups"] = [group for group in hashes.values() if len(group) > 1]
    report["image_dimensions"] = dict(collections.Counter(f'{x["width"]}x{x["height"]}' for x in report["images"]))
    report["image_total_bytes"] = sum(x["bytes"] for x in report["images"])
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 18)
    for start in range(0, len(images), 20):
        sheet = Image.new("RGB", (2000, 1160), "#edf0f3")
        draw = ImageDraw.Draw(sheet)
        for offset, path in enumerate(images[start:start + 20]):
            left, top = (offset % 5) * 400, (offset // 5) * 290
            with Image.open(path) as im:
                thumb = im.convert("RGB")
                thumb.thumbnail((390, 254))
                sheet.paste(thumb, (left + 5, top + 30))
            draw.text((left + 7, top + 5), path.name, font=font, fill="#202d39")
        sheet.save(scratch / f"contact-{start // 20 + 1}.jpg", quality=90)
    for archive in data.rglob("*.zip"):
        with zipfile.ZipFile(archive) as zip_file:
            infos = [i for i in zip_file.infolist() if not i.is_dir()]
            report["archives"].append({"name": archive.name, "bytes": archive.stat().st_size,
                "files": len(infos), "uncompressed_bytes": sum(i.file_size for i in infos),
                "extensions": dict(collections.Counter(Path(i.filename).suffix for i in infos)),
                "members": [{"name": i.filename, "bytes": i.file_size} for i in infos]})
    for workbook in data.rglob("*.xlsx"):
        wb = load_workbook(workbook, read_only=False, data_only=False)
        book = {"name": workbook.name, "sha256": digest(workbook), "sheets": []}
        for sheet in wb.worksheets:
            rows = []
            formula_cells, date_cells, date_cell_details = [], [], []
            for row in sheet.iter_rows():
                values = {}
                for cell in row:
                    if cell.value is not None:
                        values[cell.coordinate] = serial(cell.value)
                        if cell.data_type == "f":
                            formula_cells.append(cell.coordinate)
                        if cell.is_date:
                            date_cells.append(cell.coordinate)
                            date_cell_details.append({'cell': cell.coordinate, 'value': serial(cell.value),
                                'data_type': cell.data_type, 'number_format': cell.number_format})
                if values:
                    rows.append({"row": row[0].row, "cells": values})
            book["sheets"].append({"name": sheet.title, "state": sheet.sheet_state, "max_row": sheet.max_row,
                "max_column": sheet.max_column, "merged": [str(r) for r in sheet.merged_cells.ranges],
                "formula_cells": formula_cells, "date_cells": date_cells, 'date_cell_details': date_cell_details, "rows": rows})
        report["workbooks"].append(book)
        wb.close()
    pdf = next(source.glob("7. Департамент*.pdf"))
    reader = PdfReader(pdf)
    pages = [page.extract_text() for page in reader.pages]
    report["pdf"] = {"name": pdf.name, "pages": len(pages), "bytes": pdf.stat().st_size, "sha256": digest(pdf)}
    (scratch / "tz-text.txt").write_text("\n\n".join(f"PAGE {i + 1}\n{page}" for i, page in enumerate(pages)), encoding="utf-8")
    (target / "input-audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    summary = {key: report[key] for key in ["audit_date", "pdf", "image_dimensions", "image_total_bytes", "exact_duplicate_groups"]}
    summary["image_count"] = len(images)
    summary["metadata_keys"] = sorted({k for x in report["images"] for k in x["metadata_keys"]})
    summary["exif_keys"] = sorted({k for x in report["images"] for k in x["exif_keys"]})
    summary["archives"] = [{k: v for k, v in a.items() if k != "members"} for a in report["archives"]]
    summary["workbooks"] = [{**b, "sheets": [{**s, "rows": s["rows"][:20]} for s in b["sheets"]]} for b in report["workbooks"]]
    print(json.dumps(summary, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
