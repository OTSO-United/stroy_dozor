"""Compact export of confirmed event intervals and equipment counts."""

import csv
import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .catalog import class_map


def safe_text(value):
    if value is None:
        return ""
    text = str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def build_report(
    project,
    plan,
    sources,
    assessments,
    observations,
    alerts,
    format,
    plans_by_id=None,
    from_time=None,
    to_time=None,
):
    source_names = {source.id: source.name for source in sources}
    plans_by_id = plans_by_id or (
        {plan.id: plan} if plan and getattr(plan, "id", None) else {}
    )
    relevant = [
        alert
        for alert in sorted(alerts, key=lambda item: item.first_seen)
        if alert.status != "unknown"
        and (alert.details or {}).get("message")
        and not (alert.details or {}).get("suppressed")
        and (alert.details or {}).get("confirmed") is not False
    ]
    class_ids = sorted(
        {
            str(class_id)
            for alert in relevant
            for class_id in (
                set((alert.details or {}).get("expected", {}))
                | set((alert.details or {}).get("counts", {}))
            )
        },
        key=lambda value: int(value) if value.isdigit() else value,
    )
    names = class_map()
    headers = [
        "Объект",
        "Начало периода отчёта, МСК",
        "Конец периода отчёта, МСК",
        "Индекс этапа",
        "Этап",
        "Начало этапа, МСК",
        "Конец этапа, МСК",
        "Начало события, МСК",
        "Конец события, МСК",
        "Последнее наблюдение события, МСК",
        "Уведомление, МСК",
        "Источник",
        "Событие",
    ]
    for class_id in class_ids:
        name = names.get(class_id, {}).get("name", f"Класс {class_id}")
        headers.extend((f"{name}: план", f"{name}: на кадре"))

    def moscow(value):
        if not value:
            return ""
        if isinstance(value, str):
            value = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if value.tzinfo is None:
            return ""
        return value.astimezone(ZoneInfo("Europe/Moscow")).strftime("%d.%m.%Y %H:%M:%S")

    rows = []
    for alert in relevant:
        details = alert.details or {}
        historical_plan = plans_by_id.get(details.get("plan_id"))
        works = (
            {work["id"]: work for work in historical_plan.works}
            if historical_plan
            else {}
        )
        work = works.get(alert.work_id, {})
        expected = details.get("expected", {})
        counts = details.get("counts", {})
        row = [
            project.name,
            moscow(from_time),
            moscow(to_time),
            details.get("code") or work.get("code", ""),
            details.get("title") or work.get("title", ""),
            moscow(details.get("work_starts_at") or work.get("starts_at")),
            moscow(details.get("work_ends_at") or work.get("ends_at")),
            moscow(alert.first_seen),
            moscow(alert.last_seen) if alert.status == "resolved" else "",
            moscow(alert.last_seen),
            moscow(details.get("notified_at")),
            source_names.get(alert.source_id, ""),
            details.get("message", ""),
        ]
        for class_id in class_ids:
            row.extend((expected.get(class_id, ""), counts.get(class_id, "")))
        rows.append(row)

    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter=";")
        writer.writerow(headers)
        for row in rows:
            writer.writerow([safe_text(value) for value in row])
        return output.getvalue().encode("utf-8-sig")

    book = Workbook()
    sheet = book.active
    sheet.title = "События"
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [safe_text(value) if isinstance(value, str) else value for value in row]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    period = book.create_sheet("Период")
    period.append(["Объект", safe_text(project.name)])
    period.append(["Начало периода отчёта, МСК", moscow(from_time)])
    period.append(["Конец периода отчёта, МСК", moscow(to_time)])
    period.column_dimensions["A"].width = 34
    period.column_dimensions["B"].width = 48
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163E37")
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(
            62, max(18, len(header) + 4)
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
