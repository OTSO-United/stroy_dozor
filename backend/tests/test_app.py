from datetime import datetime, timedelta
from io import BytesIO
from uuid import uuid4
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from app import models as m, schemas, worker, monitor
from app.catalog import catalog
from app.planning import shift_plan
from app.rules import evaluate, severity
from app.jobs import claim, fenced, LeaseLost
from app.db import now


def post(client, url, data=None, key=None, **kwargs):
    return client.post(
        "/api/v1" + url,
        json=data,
        headers={"Idempotency-Key": key or str(uuid4())},
        **kwargs,
    )


def work(
    code="12.2",
    start="2026-09-16T09:00:00+03:00",
    end="2026-09-16T12:00:00+03:00",
    **kw,
):
    return {
        "id": str(uuid4()),
        "code": code,
        "title": "Работа " + code,
        "starts_at": start,
        "ends_at": end,
        "resources": {"0": 1, "1": 2, "16": 0},
        **kw,
    }


def approved(client, p, works):
    r = post(
        client,
        f"/projects/{p['id']}/plans",
        {"works": works, "base_plan_id": p.get("current_plan_id")},
    )
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "approved"
    return r.json()


def test_catalog_is_complete_and_normalized():
    rows = catalog()["works"]
    assert len(rows) == 377 and len({r["code"] for r in rows}) == 377
    assert all(not row["code"].endswith(".") for row in rows)


def test_idempotency_and_payload_conflict(client):
    a = post(client, "/projects", {"name": "A", "class_ids": [0, 1]}, "same")
    b = post(client, "/projects", {"name": "A", "class_ids": [0, 1]}, "same")
    conflict = post(client, "/projects", {"name": "B", "class_ids": [0, 1]}, "same")
    assert a.status_code == 201 and a.json()["id"] == b.json()["id"]
    assert conflict.status_code == 409
    assert len(client.get("/api/v1/projects").json()) == 1


def test_plan_requires_all_zero_columns(client, project):
    w = work(resources={"0": 1, "1": 2})
    r = post(client, f"/projects/{project['id']}/plans", {"works": [w]})
    assert r.status_code == 422
    w["resources"]["16"] = 0
    plan = approved(client, project, [w])
    assert plan["works"][0]["resources"]["16"] == 0


def test_no_sequence_warning_manual_completion_and_audit(client, project):
    a, b = work("12.2"), work("12.3")
    approved(client, project, [a, b])
    result = post(
        client,
        f"/projects/{project['id']}/works/{b['id']}/status",
        {"status": "in_progress"},
    )
    assert result.status_code == 200
    assert result.json()["warnings"] == []
    post(
        client,
        f"/projects/{project['id']}/works/{a['id']}/status",
        {"status": "completed"},
    )
    result = post(
        client,
        f"/projects/{project['id']}/works/{b['id']}/status",
        {"status": "in_progress"},
    )
    assert not any("12.2 Работа" in w for w in result.json()["warnings"])
    events = client.get(f"/api/v1/projects/{project['id']}/audit").json()
    assert any(e["action"] == "work.status_changed" for e in events)


def test_shift_only_conflicts_and_preserves_versions(client, project):
    a = work("12.1", end="2026-09-16T10:00:00+03:00")
    b = work("12.2", start="2026-09-16T11:00:00+03:00", end="2026-09-16T13:00:00+03:00")
    c = work("13.1")
    plan = approved(client, project, [a, b, c])
    r = post(
        client,
        f"/plans/{plan['id']}/shift-preview",
        {"work_id": a["id"], "hours": 2, "cascade": True},
    )
    assert len(r.json()["changes"]) == 2
    changed = post(
        client,
        f"/projects/{project['id']}/shift",
        {
            "work_id": a["id"],
            "hours": 2,
            "cascade": True,
            "expected_plan_id": plan["id"],
        },
    )
    assert changed.status_code == 200
    assert changed.json()["works"][1]["starts_at"].startswith("2026-09-16T12:")
    assert changed.json()["works"][2] == plan["works"][2]
    conflict = post(
        client,
        f"/projects/{project['id']}/shift",
        {"work_id": a["id"], "hours": 1, "expected_plan_id": plan["id"]},
    )
    assert conflict.status_code == 409
    versions = client.get(f"/api/v1/projects/{project['id']}/plans").json()["plans"]
    assert versions[1]["works"][0]["ends_at"] == a["ends_at"]


def test_end_cannot_precede_start():
    w = work()
    with pytest.raises(ValueError):
        shift_plan([w], w["id"], -10, True)


def test_cascade_does_not_repair_unrelated_existing_conflicts():
    a, b, c, d = work("12.1"), work("12.2"), work("13.1"), work("13.2")
    result, changes = shift_plan([a, b, c, d], a["id"], 2, True)
    assert result[2:] == [c, d]
    assert {x["work_id"] for x in changes} == {a["id"], b["id"]}


def test_csv_empty_and_invalid_timezone_return_validation(client, project):
    response = client.post(
        f"/api/v1/projects/{project['id']}/plan-imports",
        files={"file": ("empty.csv", b"")},
        headers={"Idempotency-Key": "empty"},
    )
    assert response.status_code == 422
    assert (
        post(client, "/projects", {"name": "X", "timezone": "invalid/zone"}).status_code
        == 422
    )


def test_csv_roundtrip_preserves_all_zero_columns(client, project):
    template = client.get(
        f"/api/v1/projects/{project['id']}/plan-template"
    ).content.decode("utf-8-sig")
    data = template + "12.1;Test;16.09.2026 09:00;16.09.2026 12:00;1;0;0\r\n"
    response = client.post(
        f"/api/v1/projects/{project['id']}/plan-imports",
        files={"file": ("plan.csv", data.encode("utf-8-sig"))},
        headers={"Idempotency-Key": "csv"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["works"][0]["resources"] == {"0": 1, "1": 0, "16": 0}


def test_upload_probe_model_absence_and_idempotency(client, project, runtime):
    import cv2
    import numpy as np

    _, bytes_ = cv2.imencode(".png", np.full((48, 64, 3), 128, np.uint8))
    kwargs = {
        "files": {"file": ("frame.png", bytes_.tobytes(), "image/png")},
        "data": {"name": "Кадр"},
        "headers": {"Idempotency-Key": "upload"},
    }
    r = client.post(f"/api/v1/projects/{project['id']}/media", **kwargs)
    assert r.status_code == 202, r.text
    again = client.post(f"/api/v1/projects/{project['id']}/media", **kwargs)
    assert r.json()["source"]["id"] == again.json()["source"]["id"]
    factory, path = runtime
    assert len(list((path / "media").iterdir())) == 1
    with factory() as db:
        job = claim(db)
    worker.process(job)
    source_id = r.json()["source"]["id"]
    assert client.get(f"/api/v1/sources/{source_id}/preview").status_code == 200
    analysis = post(client, f"/sources/{source_id}/analyze")
    assert (
        analysis.status_code == 409
        and analysis.json()["detail"]["code"] == "model_unavailable"
    )
    assert client.get(f"/api/v1/sources/{source_id}/observations").json() == []


def test_corrupt_file_fails_without_fake_observation(client, project, runtime):
    r = client.post(
        f"/api/v1/projects/{project['id']}/media",
        files={"file": ("bad.png", b"not-an-image")},
        headers={"Idempotency-Key": "bad"},
    )
    factory, _ = runtime
    with factory() as db:
        job = claim(db)
    worker.process(job)
    sources = client.get(f"/api/v1/projects/{project['id']}/sources").json()
    assert sources[0]["status"] == "failed"
    assert (
        client.get(f"/api/v1/sources/{r.json()['source']['id']}/observations").json()
        == []
    )


def test_stale_lease_cannot_commit(runtime, client, project):
    factory, _ = runtime
    with factory() as db:
        source = m.Source(project_id=project["id"], name="X", kind="video")
        db.add(source)
        db.flush()
        db.add(m.Job(project_id=project["id"], source_id=source.id, kind="probe"))
        db.commit()
        first = claim(db)
        token = first.token
        first.lease_until = now() - timedelta(seconds=1)
        db.commit()
        second = claim(db)
        assert token != second.token
        with pytest.raises(LeaseLost):
            fenced(db, first.id, token, status="succeeded")
        db.rollback()
        fenced(db, second.id, second.token, status="succeeded")
        db.commit()


def test_zone_cannot_reference_other_project_work(client, project, runtime):
    approved(client, project, [work()])
    factory, _ = runtime
    with factory() as db:
        source = m.Source(project_id=project["id"], name="X", kind="video")
        db.add(source)
        db.commit()
        source_id = source.id
    region = {
        "name": "Зона",
        "work_ids": [str(uuid4())],
        "polygon": [[0, 0], [1, 0], [1, 1], [0, 1]],
    }
    r = post(client, f"/sources/{source_id}/bindings", {"regions": [region]})
    assert r.status_code == 422


def test_bad_time_and_missing_idempotency(client):
    assert client.post("/api/v1/projects", json={"name": "x"}).status_code == 422
    with pytest.raises(ValueError):
        schemas.Work(**work(starts_at="2026-09-16T09:00:00"))


def test_rules_missing_excess_unknown_parallel_and_stationary():
    a = work(resources={"0": 1, "1": 1, "16": 0})
    region = {
        "id": "r",
        "name": "Зона",
        "work_ids": [a["id"]],
        "polygon": [[0, 0], [1, 0], [1, 1], [0, 1]],
        "primary": True,
        "visibility_confirmed": True,
    }
    detections = [
        {
            "class_id": "16",
            "bbox": [0.1, 0.1, 0.2, 0.2],
            "activity": "unknown",
            "motion": "stationary_observed",
        }
    ]
    t = datetime.fromisoformat(a["starts_at"])
    result = evaluate([a], [region], detections, t, ["0", "1", "16"], True)[0]
    assert {f["kind"] for f in result["findings"]} == {"missing", "excess"}
    assert (
        next(f for f in result["findings"] if f["kind"] == "excess")["message"]
        == "Сверх плана 1 ед. техники"
    )
    duplicate = {**detections[0], "bbox": [0.3, 0.1, 0.4, 0.2]}
    doubled = evaluate(
        [a], [region], [*detections, duplicate], t, ["0", "1", "16"], True
    )[0]
    assert (
        next(f for f in doubled["findings"] if f["kind"] == "excess")["message"]
        == "Сверх плана 2 ед. техники"
    )
    assert all(f["kind"] != "idle" for f in result["findings"])
    detections[0]["activity"] = "idle"
    idle_result = evaluate([a], [region], detections, t, ["0", "1", "16"], True)[0]
    assert next(f for f in idle_result["findings"] if f["kind"] == "idle")[
        "classes"
    ] == [{"class_id": "16", "count": 1}]
    unknown = evaluate([a], [region], detections, t, ["0", "1", "16"], False)[0]
    assert (
        all(v is None for v in unknown["counts"].values()) and unknown["findings"] == []
    )
    b = work("13.1", resources={"0": 0, "1": 0, "16": 1})
    region["work_ids"].append(b["id"])
    result = evaluate([a, b], [region], detections, t, ["0", "1", "16"], True)[0]
    assert "excess" not in {f["kind"] for f in result["findings"]}


def test_severity_is_strict_and_requires_both_conditions():
    settings = schemas.Settings().model_dump()
    assert (
        severity({"kind": "missing", "missing_ratio": 0.5}, 4000, settings) == "warning"
    )
    assert (
        severity({"kind": "missing", "missing_ratio": 0.9}, 3599, settings) == "warning"
    )
    assert (
        severity({"kind": "missing", "missing_ratio": 0.9}, 3600, settings)
        == "critical"
    )


def test_report_exports_only_significant_event_intervals(client, project):
    approved(client, project, [work(title="=2+2")])
    r = client.get(f"/api/v1/projects/{project['id']}/reports/xlsx")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["События", "Период"]
    assert wb["События"].max_row == 1
    r = client.get(f"/api/v1/projects/{project['id']}/reports/csv")
    assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf")


def test_monitor_critical_duration_resets_when_ratio_falls(runtime, client, project):
    factory, _ = runtime
    w = work(resources={"0": 0, "1": 4, "16": 0})
    plan = approved(client, project, [w])
    with factory() as db:
        source = m.Source(project_id=project["id"], name="X", kind="video")
        db.add(source)
        db.flush()
        binding = m.Binding(
            project_id=project["id"],
            source_id=source.id,
            revision=1,
            regions=[
                {
                    "id": "r",
                    "name": "Зона",
                    "work_ids": [w["id"]],
                    "polygon": [[0, 0], [1, 0], [1, 1], [0, 1]],
                    "primary": True,
                    "visibility_confirmed": True,
                }
            ],
        )
        db.add(binding)
        db.flush()
        settings = {
            **schemas.Settings().model_dump(),
            "critical_after_seconds": 2,
            "absence_confirm_seconds": 0,
        }
        job = m.Job(
            project_id=project["id"],
            source_id=source.id,
            kind="analyze",
            payload={
                "plan_id": plan["id"],
                "binding_id": binding.id,
                "settings": settings,
            },
        )
        db.add(job)
        db.flush()
        base = datetime.fromisoformat(w["starts_at"])
        for i, count in enumerate([0, 0, 3, 0, 0, 0]):
            o = m.Observation(
                project_id=project["id"],
                source_id=source.id,
                run_id=job.id,
                sample_index=i,
                captured_at=base + timedelta(seconds=i),
                offset_seconds=i,
                detections=[{"class_id": "1", "bbox": [0.1, 0.1, 0.2, 0.2]}] * count,
                quality={"usable": True},
                model={
                    "supported_classes": ["0", "1", "16"],
                    "sha256": "synthetic-test",
                },
                evidence_key="test-only",
            )
            db.add(o)
            db.flush()
            monitor.assess_one(db, o)
            db.commit()
            alert = db.scalar(select(m.Alert).where(m.Alert.status == "open"))
            assert alert.severity == ("critical" if i == 5 else "warning")
        # A long missing interval starts a new episode rather than extending the old one.
        o = m.Observation(
            project_id=project["id"],
            source_id=source.id,
            run_id=job.id,
            sample_index=6,
            captured_at=base + timedelta(seconds=100),
            offset_seconds=100,
            detections=[],
            quality={"usable": True},
            model={"supported_classes": ["0", "1", "16"]},
            evidence_key="test-only",
        )
        db.add(o)
        db.flush()
        monitor.assess_one(db, o)
        db.commit()
        alerts = list(db.scalars(select(m.Alert)))
        assert len(alerts) == 2
        latest = max(alerts, key=lambda alert: alert.details["episode"])
        assert latest.duration_seconds == 0
        assert latest.details["episode"] == 2
        assert sum(alert.status == "interrupted" for alert in alerts) == 1


def test_idle_alert_has_class_scoped_keyframe_evidence(runtime, client, project):
    factory, _ = runtime
    stage = work(resources={"0": 0, "1": 1, "16": 0})
    plan = approved(client, project, [stage])
    with factory() as db:
        source = m.Source(project_id=project["id"], name="Камера простоя", kind="rtsp")
        db.add(source)
        db.flush()
        binding = m.Binding(
            project_id=project["id"],
            source_id=source.id,
            revision=1,
            regions=[
                {
                    "id": "zone",
                    "name": "Весь кадр",
                    "work_ids": [stage["id"]],
                    "polygon": [[0, 0], [1, 0], [1, 1], [0, 1]],
                    "primary": True,
                    "visibility_confirmed": True,
                }
            ],
        )
        db.add(binding)
        db.flush()
        job = m.Job(
            project_id=project["id"],
            source_id=source.id,
            kind="analyze",
            payload={
                "plan_id": plan["id"],
                "binding_id": binding.id,
                "settings": schemas.Settings().model_dump(),
                "sample_seconds": 60,
            },
        )
        db.add(job)
        db.flush()
        observation = m.Observation(
            project_id=project["id"],
            source_id=source.id,
            run_id=job.id,
            sample_index=0,
            captured_at=datetime.fromisoformat(stage["starts_at"]),
            offset_seconds=0,
            detections=[
                {
                    "class_id": "1",
                    "bbox": [0.1, 0.1, 0.3, 0.4],
                    "confidence": 0.9,
                    "observed": True,
                    "activity": "idle",
                }
            ],
            quality={"usable": True},
            model={"supported_classes": ["0", "1", "16"], "sha256": "test"},
            evidence_key="test-only",
        )
        db.add(observation)
        db.flush()
        monitor.assess_one(db, observation)
        db.commit()
        alert = db.scalar(select(m.Alert).where(m.Alert.kind == "idle"))
        assert alert is not None and alert.status == "open"
        alert_id = alert.id
    response = client.get(f"/api/v1/alerts/{alert_id}/evidence?class_id=1")
    assert response.status_code == 200
    evidence = response.json()
    assert evidence["deviation_class_ids"] == ["1"]
    assert evidence["missing_class_ids"] == []
    assert evidence["keyframes"][0]["detections"][0]["activity"] == "idle"


def test_monitor_confirms_absence_over_time_without_duplicate_signals(
    runtime, client, project
):
    factory, _ = runtime
    stage = work(resources={"0": 0, "1": 1, "16": 0})
    plan = approved(client, project, [stage])
    with factory() as db:
        source = m.Source(project_id=project["id"], name="Сглаживание", kind="video")
        db.add(source)
        db.flush()
        binding = m.Binding(
            project_id=project["id"],
            source_id=source.id,
            revision=1,
            regions=[
                {
                    "id": "r",
                    "name": "Весь кадр",
                    "work_ids": [stage["id"]],
                    "polygon": [[0, 0], [1, 0], [1, 1], [0, 1]],
                    "primary": True,
                    "visibility_confirmed": True,
                }
            ],
        )
        db.add(binding)
        db.flush()
        settings = {
            **schemas.Settings().model_dump(),
            "absence_confirm_seconds": 2,
        }
        job = m.Job(
            project_id=project["id"],
            source_id=source.id,
            kind="analyze",
            payload={
                "plan_id": plan["id"],
                "binding_id": binding.id,
                "settings": settings,
                "sample_seconds": 1,
            },
        )
        db.add(job)
        db.flush()
        base = datetime.fromisoformat(stage["starts_at"])
        for index, present in enumerate([False, True, False, False, False]):
            observation = m.Observation(
                project_id=project["id"],
                source_id=source.id,
                run_id=job.id,
                sample_index=index,
                captured_at=base + timedelta(seconds=index),
                offset_seconds=index,
                detections=(
                    [{"class_id": "1", "bbox": [0.1, 0.1, 0.2, 0.2]}] if present else []
                ),
                quality={"usable": True},
                model={"supported_classes": ["0", "1", "16"]},
                evidence_key="test-only",
            )
            db.add(observation)
            db.flush()
            monitor.assess_one(db, observation)
            db.commit()

        alerts = list(db.scalars(select(m.Alert)))
        assert len(alerts) == 2
        alert = max(alerts, key=lambda item: item.details["episode"])
        assert alert.status == "open"
        assert alert.duration_seconds == 2
        assert alert.details["confirmed"] is True
        assert alert.details["episode"] == 2
        assert alert.details["confirmed_absence_seconds"] == 0
