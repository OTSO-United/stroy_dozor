"""Verify the real HTTP XLSX downloaded by the isolated APP-29 browser check."""

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


path = Path(__file__).resolve().parents[1] / "tmp/app29-browser/report.xlsx"
book = load_workbook(path)
assert book.sheetnames == ["События", "Период"]
headers, *values = list(book["События"].values)
rows = [dict(zip(headers, row)) for row in values]
assert rows
assert {row["Индекс этапа"] for row in rows} == {"12.1", "12.2"}
assert all(row["Объект"] == "APP-29 · Изолированная проверка" for row in rows)
assert all(row["Этап"] and row["Самосвал: план"] == 4 for row in rows)
assert all(
    row["Экскаватор: план"] == 1 and row["Экскаватор: на кадре"] == 1 for row in rows
)
assert any(row["Самосвал: на кадре"] == 0 for row in rows)
assert all(
    row["Начало периода отчёта, МСК"] == book["Период"]["B2"].value for row in rows
)
assert all(
    row["Конец периода отчёта, МСК"] == book["Период"]["B3"].value for row in rows
)
for row in rows:
    dates = {
        header: datetime.strptime(row[header], "%d.%m.%Y %H:%M:%S")
        for header in headers
        if header.endswith(", МСК") and row[header]
    }
    assert (
        dates["Начало этапа, МСК"]
        <= dates["Начало события, МСК"]
        < dates["Конец этапа, МСК"]
    )
    assert dates["Начало события, МСК"] <= dates["Последнее наблюдение события, МСК"]
    assert dates["Начало события, МСК"] < dates["Конец периода отчёта, МСК"]
assert any(row["Конец события, МСК"] is None for row in rows)
print(
    json.dumps(
        {
            "xlsx": str(path),
            "rows": len(rows),
            "columns": len(headers),
            "stages": ["12.1", "12.2"],
            "dates_and_values": "passed",
        },
        ensure_ascii=True,
    )
)
